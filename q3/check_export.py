"""Independent saved-workbook readback; missing/NaN/failed evidence is rejected."""
import argparse
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from q2.archive import write_json
from q2.inputs import sha256
from q2.provenance import portable_artifact_sha256
from q3.validation import verified
from q3.export import rounded,delivery_sources


def compare_numeric(actual,expected):
    if any(v is None or isinstance(v,(str,bool)) for row in actual for v in row):
        raise ValueError("Missing or nonnumeric Excel result")
    values=np.asarray(actual,dtype=float)
    if values.shape != expected.shape or not np.isfinite(values).all() or not np.isfinite(expected).all():
        raise ValueError("Wrong shape or nonfinite workbook data")
    diff=float(np.max(np.abs(values-expected)))
    if not np.isfinite(diff) or diff != 0:
        raise ValueError(f"Workbook/source mismatch: {diff}")
    return diff


def check(directory="results/q3",workbook="results/result3.xlsx"):
    directory=Path(directory)
    write_json(directory/"export_verification.json",{"passed":False,"status":"checking"})
    v,r,f=verified(directory)
    # Artifact Tool omits the optional worksheet dimension metadata, so
    # openpyxl's read-only mode reports max_row/max_column as None. This
    # workbook is small enough for a normal independent load.
    wb=load_workbook(workbook,read_only=False,data_only=True)
    try:
        if wb.sheetnames != ["Sheet1"]:
            raise ValueError("Wrong workbook sheets")
        sheet=wb.active
        expected=np.c_[f["time_s"][1:],rounded(f["moisture"][1:])]
        if sheet.max_row != len(expected)+1 or sheet.max_column != 22:
            raise ValueError("Unexpected workbook dimensions")
        iterator=sheet.iter_rows()
        header=next(iterator)
        if header[0].value != r["identity"]["inputs"]["q3_template"]["A1"]:
            raise ValueError("Wrong template A1")
        np.testing.assert_array_equal([c.value for c in header[1:]],np.arange(21)/10)
        actual=[]
        for row in iterator:
            if any(c.number_format != "0.0000" for c in row[1:]):
                raise ValueError("A moisture cell does not display four decimal places")
            actual.append([c.value for c in row])
        # XLSX decimal serialization can change an endpoint by one ULP. Compare
        # exact regular seconds and sub-nanosecond terminal time independently.
        actual_array=np.array(actual,dtype=object)
        difference=compare_numeric(actual_array[:,1:].tolist(),expected[:,1:])
        times=np.array(actual_array[:,0],dtype=float)
        if not np.isfinite(times).all() or np.any(np.diff(times)<=0):
            raise ValueError("Invalid or duplicate time axis")
        np.testing.assert_array_equal(times[:-1],expected[:-1,0])
        time_error=abs(times[-1]-expected[-1,0])
        if time_error>1e-9:
            raise ValueError("Incorrect terminal time")
    finally:
        wb.close()
    table=np.loadtxt(directory/"table5.csv",delimiter=",",skiprows=1)
    indices=np.searchsorted(f["time_s"],table[:,0]*3600)
    # Multiplication h->s can round to either adjacent float; nearest recorded row.
    indices=np.array([np.argmin(abs(f["time_s"]-t*3600)) for t in table[:,0]])
    np.testing.assert_allclose(f["time_s"][indices],table[:,0]*3600,rtol=0,atol=1e-9)
    np.testing.assert_array_equal(table[:,1:],f["moisture"][indices][:,[0,5,10,15,20]])
    result={"passed":True,"cells_checked":len(expected)*21,"rows":len(expected),"columns":22,
            "moisture_max_abs_difference":difference,"terminal_time_difference_s":float(time_error),
            "workbook_sha256":sha256(workbook),"verification_sha256":portable_artifact_sha256(directory/"verification.json"),
            "delivery_sources":delivery_sources(),"table5_sha256":portable_artifact_sha256(directory/"table5.csv")}
    write_json(directory/"export_verification.json",result)
    return result


if __name__ == "__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("--directory",default="results/q3")
    parser.add_argument("--workbook",default="results/result3.xlsx")
    a=parser.parse_args()
    print(check(a.directory,a.workbook))

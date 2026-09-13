"""Reproducible cases. Every new run starts from the uniform original state."""
import argparse
from concurrent.futures import ProcessPoolExecutor, as_completed
import copy
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from common.hashing import file_record
from q1.fvm import Grid
from q2.inputs import read_config, read_inputs, sha256
from q2.model import CoupledRadialModel
from q2.archive import write_json
from q3.solver import integrate_event, TRACE_COLUMNS
from q3.provenance import snapshot, load_run


def cases(config, include_refinement=True):
    extra_names = {case["name"] for case in config["refinement_cases"]}
    result = [{"name": f"base_N{n}", "N": n} for n in config["grids"]
              if f"base_N{n}" not in extra_names]
    result += copy.deepcopy(config["temporal_cases"])
    for mode in ["terminal_hold", "nominal"]:
        result.append({"name": mode, "N": config["sensitivity_grid"], "mode": mode})
    for key in ["hm", "D_prefactor"]:
        for factor in [0.9, 1.1]:
            result.append({"name": f"{key}_{factor:.1f}", "N": config["sensitivity_grid"],
                           "parameter": key, "factor": factor})
    all_cases = result + copy.deepcopy(config["refinement_cases"])
    names = [case["name"] for case in all_cases]
    if len(names) != len(set(names)):
        raise ValueError("Duplicate configured case name")
    if config["formal_case"] not in names:
        raise ValueError("formal_case is not a configured case")
    if any(name not in names for pair in config["temporal_comparisons"] for name in pair):
        raise ValueError("Unknown temporal comparison case")
    return all_cases if include_refinement else result


def compute(data_root, directory, case):
    directory = Path(directory)/"runs"/case["name"]
    config = read_config("configs/q3.json")
    inherited = read_config(config["q2_config"])
    env, inputs = read_inputs(data_root, inherited, mode=case.get("mode"))
    inputs["source_names"] = {k: v.replace("\\", "/") for k, v in inputs["source_names"].items()}
    template = Path(data_root)/"附件"/"附件3"/"result3.xlsx"
    wb = load_workbook(template, read_only=True)
    try:
        if wb.sheetnames != ["Sheet1"]:
            raise ValueError("Unexpected result3 template")
        inputs["q3_template"] = {"sha256": sha256(template), "A1": wb.active["A1"].value}
    finally:
        wb.close()
    settings = copy.deepcopy(inherited["solver"])
    if case.get("tight"):
        for key in ["rtol", "atol_T", "atol_C"]:
            settings[key] *= 0.1
    if case.get("half_step"):
        for key in ["max_step_observed", "max_step_extended"]:
            settings[key] *= 0.5
    parameters = dict(inherited["parameters"])
    if "parameter" in case:
        parameters[case["parameter"]] *= case["factor"]
    identity = {"case": case, "inputs": inputs, "parameters": parameters,
                "settings": settings, "event": config}
    if (directory/"run.json").exists():
        record, _ = load_run(directory)
        if record["identity"] != identity:
            raise ValueError("Existing run configuration/input mismatch; use a new directory")
        return case["name"], record["root"]["time_h"], "verified cache"
    directory.mkdir(parents=True, exist_ok=True)
    source = snapshot()
    write_json(directory/"run.json", {"status": "running", "identity": identity, "source": source})
    try:
        grid = Grid(case["N"], parameters["R"], parameters["L"])
        solution = integrate_event(CoupledRadialModel(grid, parameters, env), settings, config)
        np.savez_compressed(directory/"fields.npz", **{k:v for k,v in solution.items()
                                                     if isinstance(v, np.ndarray) and k != "trace"})
        np.savetxt(directory/"accepted_steps.csv", solution["trace"], delimiter=",",
                   header=",".join(TRACE_COLUMNS), comments="", fmt="%.17g")
        write_json(directory/"run.json", {"status": "computed", "identity": identity,
            "source": source, "root": solution["root"], "diagnostics": solution["diagnostics"],
            "files": {p: file_record(directory / p) for p in ["fields.npz", "accepted_steps.csv"]}})
        return case["name"], solution["root"]["time_h"], solution["diagnostics"]["wall_seconds"]
    except Exception as exc:
        write_json(directory/"failure.json", {"error": repr(exc), "identity": identity, "source": source})
        raise RuntimeError(f"{case['name']}: {exc!r}") from exc


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-root", required=True)
    parser.add_argument("--directory", default="results/q3")
    parser.add_argument("--cases", nargs="*")
    parser.add_argument("--workers", type=int, default=1)
    args = parser.parse_args()
    selected = cases(read_config("configs/q3.json"), include_refinement=bool(args.cases))
    if args.cases:
        selected = [c for c in selected if c["name"] in args.cases]
        if len(selected) != len(set(args.cases)):
            raise ValueError("Unknown case")
    with ProcessPoolExecutor(max_workers=args.workers) as pool:
        jobs = [pool.submit(compute, args.data_root, args.directory, case) for case in selected]
        for job in as_completed(jobs):
            print(json.dumps(job.result()), flush=True)


if __name__ == "__main__":
    main()

"""Validated Q2 inputs and explicit long-horizon environment extensions."""
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


def sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_config(path="configs/q2.json"):
    return json.loads(Path(path).read_text(encoding="utf-8"))


class LongEnvironment:
    MODES = {"last_hour_mean", "terminal_hold", "nominal"}

    def __init__(self, observations, horizon_s, mode="last_hour_mean", mean_start_s=10800):
        values = np.asarray(observations, dtype=float)
        expected_times = np.arange(0, 14401, 60, dtype=float)
        if values.shape != (241, 3) or not np.all(np.isfinite(values)):
            raise ValueError("Expected 241 finite environment records")
        if not np.array_equal(values[:, 0], expected_times) or np.any(values[:, 2] <= 0):
            raise ValueError("Unexpected environment time axis or nonpositive moisture drive")
        if mode not in self.MODES or horizon_s < values[-1, 0]:
            raise ValueError("Invalid environment extension or horizon")
        self.observations = values.copy()
        self.horizon_s = float(horizon_s)
        self.mode = mode
        if mode == "last_hour_mean":
            extension = values[values[:, 0] >= mean_start_s, 1:].mean(axis=0)
        elif mode == "terminal_hold":
            extension = values[-1, 1:].copy()
        else:
            extension = np.array([50.0, 0.05])
        self.extension = extension
        self._history_min = np.minimum.accumulate(values[:, 1:], axis=0)
        self._history_max = np.maximum.accumulate(values[:, 1:], axis=0)

    def __call__(self, t):
        query = np.asarray(t, dtype=float)
        if np.any(~np.isfinite(query)) or np.any(query < 0) or np.any(query > self.horizon_s):
            raise ValueError("Environment query outside configured horizon")
        source = self.observations
        temperature = np.interp(query, source[:, 0], source[:, 1])
        moisture = np.interp(query, source[:, 0], source[:, 2])
        after = query > source[-1, 0]
        temperature = np.where(after, self.extension[0], temperature)
        moisture = np.where(after, self.extension[1], moisture)
        return temperature, moisture

    def history_extrema(self, t):
        query = np.asarray(t, dtype=float)
        temperature, moisture = self(query)
        source = self.observations
        index = np.searchsorted(source[:, 0], np.minimum(query, source[-1, 0]), side="right") - 1
        current = np.stack((temperature, moisture), axis=-1)
        low = np.minimum(self._history_min[index], current)
        high = np.maximum(self._history_max[index], current)
        return low[..., 0], high[..., 0], low[..., 1], high[..., 1]

    @property
    def breaks(self):
        return np.r_[self.observations[:, 0], self.horizon_s]


def read_inputs(data_root, config, mode=None):
    root = Path(data_root)
    paths = {
        "environment": root / "附件" / "附件1.xlsx",
        "template": root / "附件" / "附件3" / "result2.xlsx",
        "problem": root / "A题.pdf",
    }
    hashes = {key: sha256(path) for key, path in paths.items()}
    for key, expected in config["input_sha256"].items():
        if hashes[key] != expected:
            raise ValueError(f"Unreviewed {key} input hash: {hashes[key]}")
    workbook = load_workbook(paths["environment"], data_only=True, read_only=True)
    try:
        rows = list(workbook["Sheet1"].values)
    finally:
        workbook.close()
    if rows[0][:3] != ("时间", "温度", "水分浓度"):
        raise ValueError("Unexpected environment headers")
    observations = np.asarray([row[:3] for row in rows[1:]], dtype=float)
    template = load_workbook(paths["template"], read_only=True)
    try:
        if template.sheetnames != ["温度", "水分浓度"]:
            raise ValueError("Unexpected result2 template sheets")
        header = template.worksheets[0]["A1"].value
    finally:
        template.close()
    env_cfg = config["environment"]
    selected_mode = mode or env_cfg["default_extension"]
    environment = LongEnvironment(
        observations,
        config["horizon_s"],
        selected_mode,
        env_cfg["mean_window_start_s"],
    )
    metadata = {
        "sha256": hashes,
        "source_names": {key: str(path.relative_to(root)) for key, path in paths.items()},
        "template_A1": header,
        "environment_extension": {
            "mode": selected_mode,
            "temperature_C": float(environment.extension[0]),
            "moisture": float(environment.extension[1]),
            "observed_end_s": env_cfg["observed_end_s"],
        },
    }
    return environment, metadata

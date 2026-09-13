"""Stream-read result2.xlsx and compare every result cell with verified sources."""
import argparse
import json
from pathlib import Path
from pathlib import PurePosixPath
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import numpy as np
from openpyxl import load_workbook

from common.hashing import file_record
from .archive import write_json
from .export import rounded_array, verified_source
from .provenance import delivery_snapshot


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def frozen_sheet_names(workbook):
    frozen = []
    with ZipFile(workbook) as archive:
        book = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {
            item.attrib["Id"]: item.attrib["Target"]
            for item in relationships.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
        }
        sheets = book.find(f"{{{MAIN_NS}}}sheets")
        if sheets is None:
            raise ValueError("Workbook has no worksheets")
        for sheet in sheets:
            relation = sheet.attrib[f"{{{OFFICE_REL_NS}}}id"]
            target = targets[relation]
            part = target.lstrip("/") if target.startswith("/") else str(PurePosixPath("xl") / target)
            root = ET.fromstring(archive.read(part))
            pane = root.find(
                f"{{{MAIN_NS}}}sheetViews/{{{MAIN_NS}}}sheetView/{{{MAIN_NS}}}pane"
            )
            if pane is None:
                continue
            if (
                pane.attrib.get("state") == "frozen"
                and pane.attrib.get("topLeftCell") == "B2"
                and float(pane.attrib.get("xSplit", 0)) == 1.0
                and float(pane.attrib.get("ySplit", 0)) == 1.0
            ):
                frozen.append(sheet.attrib["name"])
    return frozen


def compare_chunk(rows, expected_times, expected_values, sheet_name):
    try:
        actual = np.asarray(rows, dtype=float)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Nonnumeric result value in {sheet_name}") from exc
    expected_times = np.asarray(expected_times, dtype=float)
    expected_values = np.asarray(expected_values, dtype=float)
    expected_shape = (len(expected_times), expected_values.shape[1] + 1)
    if actual.shape != expected_shape:
        raise ValueError(f"Unexpected result block shape in {sheet_name}: {actual.shape}")
    if not np.all(np.isfinite(actual)):
        raise ValueError(f"Missing or nonfinite result value in {sheet_name}")
    if not np.array_equal(actual[:, 0], expected_times):
        raise ValueError(f"Broken time axis in {sheet_name}")
    if not np.all(np.isfinite(expected_values)):
        raise ValueError(f"Nonfinite verified source value in {sheet_name}")
    delta = np.abs(actual[:, 1:] - expected_values)
    if not np.all(np.isfinite(delta)):
        raise ValueError(f"Nonfinite readback difference in {sheet_name}")
    return float(np.max(delta)), int(expected_values.size)


def failure_record(directory, workbook, exc):
    result = {
        "passed": False,
        "error_type": type(exc).__name__,
        "message": str(exc),
        "workbook": str(workbook),
    }
    if workbook.is_file():
        result.update({
            "workbook_hash": file_record(workbook),
            "workbook_bytes": workbook.stat().st_size,
        })
    write_json(directory / "export_verification.json", result)


def check(directory="results/q2", workbook="results/result2.xlsx"):
    directory, workbook = Path(directory), Path(workbook)
    book = None
    try:
        data, _, manifest, _ = verified_source(directory)
        book = load_workbook(workbook, read_only=True, data_only=True)
        max_difference = 0.0
        checked = 0
        if book.sheetnames != ["温度", "水分浓度"]:
            raise ValueError("Unexpected result2 sheet names")
        if frozen_sheet_names(workbook) != book.sheetnames:
            raise ValueError("Unexpected result2 freeze panes")
        for sheet_name, key in (("温度", "temperature_C"), ("水分浓度", "moisture")):
            sheet = book[sheet_name]
            if sheet.max_row is not None and sheet.max_row != 259201:
                raise ValueError(f"Unexpected {sheet_name} declared row count")
            if sheet.max_column is not None and sheet.max_column != 22:
                raise ValueError(f"Unexpected {sheet_name} dimensions")
            rows = sheet.iter_rows(values_only=True)
            header = next(rows)
            if header != (manifest["inputs"]["template_A1"], *[j / 10 for j in range(21)]):
                raise ValueError(f"Unexpected {sheet_name} header")
            for start in range(1, 259201, 21600):
                end = min(start + 21600, 259201)
                expected = rounded_array(data[key][start:end])
                difference, count = compare_chunk(
                    [next(rows) for _ in range(end - start)],
                    np.arange(start, end), expected, sheet_name,
                )
                max_difference = max(max_difference, difference)
                checked += count
            if next(rows, None) is not None:
                raise ValueError(f"Unexpected extra rows in {sheet_name}")
            for coordinate in ("B2", "B129601", "B259201"):
                cell = sheet[coordinate]
                if cell.data_type != "n" or cell.number_format != "0.0000":
                    raise ValueError(f"Unexpected numeric format at {sheet_name}!{coordinate}")
        result = {
            "passed": max_difference == 0.0 and checked == 10886400,
            "workbook_hash": file_record(workbook),
            "workbook_bytes": workbook.stat().st_size,
            "numeric_result_cells_checked": checked,
            "max_absolute_readback_difference": max_difference,
            "sheets": ["温度", "水分浓度"],
            "rows_per_sheet": 259201,
            "columns_per_sheet": 22,
            "numerical_verification_hash": file_record(directory / "verification.json"),
            "archive_manifest_hash": file_record(directory / "archive" / "manifest.json"),
            "delivery_source": delivery_snapshot(),
            "writer": "openpyxl write-only streaming",
        }
        write_json(directory / "export_verification.json", result)
        if not result["passed"]:
            raise ValueError("Q2 workbook readback failed")
        return result
    except Exception as exc:
        failure_record(directory, workbook, exc)
        raise
    finally:
        if book is not None:
            book.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--directory", default="results/q2")
    parser.add_argument("--workbook", default="results/result2.xlsx")
    args = parser.parse_args()
    print(json.dumps(check(args.directory, args.workbook), ensure_ascii=False))

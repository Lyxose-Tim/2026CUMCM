"""Plain-openpyxl writers shared by the four result workbooks."""
from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BODY_FONT = Font(name="Arial", size=10, color="1F2937")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")


def round_half_up(value, places=4):
    if value is None:
        return None
    number = float(value)
    if np.isnan(number):
        return None
    if not np.isfinite(number):
        raise ValueError("A nonfinite workbook value cannot be exported")
    quantum = Decimal(1).scaleb(-places)
    return float(Decimal(str(number)).quantize(quantum, rounding=ROUND_HALF_UP))


def rounded_matrix(values, places=4):
    array = np.asarray(values, dtype=float)
    result = np.empty(array.shape, dtype=object)
    for index, value in np.ndenumerate(array):
        result[index] = round_half_up(value, places)
    return result


def _configure_sheet(sheet, columns):
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 22
    sheet.column_dimensions["A"].width = 26
    for column in range(2, columns + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 12


def _write_only_cell(sheet, value, number_format, *, header=False):
    cell = WriteOnlyCell(sheet, value=value)
    cell.number_format = number_format
    cell.font = HEADER_FONT if header else BODY_FONT
    if header:
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
    return cell


def write_stream_workbook(output, sheets, *, value_format="0.0000", time_format="0"):
    """Write large flat result sheets without retaining cells in memory."""
    workbook = Workbook(write_only=True)
    for spec in sheets:
        header = spec["header"]
        sheet = workbook.create_sheet(spec["name"])
        _configure_sheet(sheet, len(header))
        sheet.append([
            _write_only_cell(sheet, value, "General", header=True) for value in header
        ])
        count = 0
        for time_s, values in zip(spec["times"], spec["values"]):
            row = [_write_only_cell(sheet, int(time_s), time_format)]
            row.extend(
                _write_only_cell(sheet, round_half_up(value), value_format)
                for value in values
            )
            sheet.append(row)
            count += 1
        if count != len(spec["times"]):
            raise ValueError(f"Workbook row source ended early for {spec['name']}")
    destination = Path(output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(destination.stem + ".tmp" + destination.suffix)
    workbook.save(temporary)
    temporary.replace(destination)
    return destination


def write_dense_workbook(output, sheets, *, value_format="0.0000", time_format="0.0000"):
    """Write moderate result workbooks in normal openpyxl mode."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for spec in sheets:
        header = spec["header"]
        sheet = workbook.create_sheet(spec["name"])
        _configure_sheet(sheet, len(header))
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
        for time_s, values in zip(spec["times"], spec["values"]):
            sheet.append([float(time_s), *[round_half_up(value) for value in values]])
        for row in sheet.iter_rows(min_row=2):
            row[0].number_format = time_format
            row[0].font = BODY_FONT
            for cell in row[1:]:
                cell.number_format = value_format
                cell.font = BODY_FONT
    destination = Path(output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(destination.stem + ".tmp" + destination.suffix)
    workbook.save(temporary)
    temporary.replace(destination)
    return destination

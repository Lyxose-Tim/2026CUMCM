"""Independent readback for result4.xlsx with out-of-material blanks."""
import argparse
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from common.hashing import file_record, verify_file
from q2.archive import write_json
from q2.inputs import sha256
from q2.provenance import portable_artifact_sha256
from q4.export import delivery_sources, rounded, workbook_headers
from q4.validation import verified


def assert_headers(header, record):
    expected = workbook_headers(record)
    if len(header) != len(expected):
        raise ValueError("Wrong result4 header width")
    if header[0] != expected[0]:
        raise ValueError("Wrong result4 A1 header")
    actual_radius = np.asarray(header[1:22], dtype=float)
    expected_radius = np.asarray(expected[1:22], dtype=float)
    if not np.allclose(actual_radius, expected_radius, rtol=0, atol=1e-12):
        raise ValueError("Wrong result4 radius headers: expected centimeters 0-2 at 0.1 cm spacing")
    if header[22] != expected[22]:
        raise ValueError("Wrong result4 dynamic surface header")


def check(directory="results/q4", workbook="results/result4.xlsx"):
    directory = Path(directory)
    write_json(directory / "export_verification.json", {"passed": False, "status": "checking"})
    verification, record, fields = verified(directory)
    radius_manifest_path = directory / "radius_input_manifest.json"
    radius_manifest = json.loads(radius_manifest_path.read_text(encoding="utf-8"))
    if (radius_manifest.get("schema_version") != 1
            or radius_manifest.get("attachment_sha256")
            != record["identity"]["inputs"]["q4_radius"]["sha256"]
            or radius_manifest.get("observed_end_s")
            != record["identity"]["inputs"]["q4_radius"]["observed_end_s"]):
        raise ValueError("Q4 radius observation evidence differs from the verified input")
    verify_file("q4/export.py", radius_manifest["generator"])
    verify_file(directory / "verification.json", radius_manifest["verification"])
    verify_file(directory / "radius_observations.csv", radius_manifest["csv"])
    wb = load_workbook(workbook, read_only=False, data_only=True)
    try:
        if wb.sheetnames != ["Sheet1"]:
            raise ValueError("Wrong Q4 workbook sheets")
        sheet = wb.active
        expected = np.c_[fields["time_s"][1:], rounded(fields["moisture"][1:])]
        if sheet.max_row != len(expected) + 1 or sheet.max_column != 23:
            raise ValueError("Unexpected Q4 workbook dimensions")
        if sheet.freeze_panes != "B2":
            raise ValueError("Missing Q4 freeze panes")
        header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        assert_headers(header, record)
        actual = []
        for row in sheet.iter_rows(min_row=2):
            if any(cell.number_format != "0.0000" for cell in row[1:] if cell.value is not None):
                raise ValueError("A Q4 moisture cell does not display four decimal places")
            actual.append([cell.value for cell in row])
    finally:
        wb.close()
    if len(actual) != len(expected):
        raise ValueError("Wrong Q4 row count")
    values = np.asarray([[np.nan if value is None else value for value in row] for row in actual], dtype=float)
    np.testing.assert_array_equal(values[:-1, 0], expected[:-1, 0])
    terminal_time_difference = abs(values[-1, 0] - expected[-1, 0])
    if terminal_time_difference > 1e-9:
        raise ValueError("Incorrect Q4 terminal time")
    actual_values = values[:, 1:]
    expected_values = expected[:, 1:]
    if not np.array_equal(np.isnan(actual_values), np.isnan(expected_values)):
        raise ValueError("Q4 workbook blank mask differs from verified source")
    diff = float(np.nanmax(np.abs(actual_values - expected_values)))
    if not np.isfinite(diff) or diff != 0:
        raise ValueError(f"Q4 workbook/source mismatch: {diff}")
    table = np.loadtxt(directory / "table6.csv", delimiter=",", skiprows=1)
    indices = np.array([np.argmin(abs(fields["time_s"] - t * 3600)) for t in table[:, 0]])
    np.testing.assert_allclose(fields["time_s"][indices], table[:, 0] * 3600, rtol=0, atol=1e-9)
    np.testing.assert_allclose(table[:, 1:], fields["moisture"][indices][:, [0, 5, 10, 21]], equal_nan=True)
    result = {
        "passed": True,
        "cells_checked": int(np.isfinite(expected_values).sum()),
        "blank_cells_checked": int(np.isnan(expected_values).sum()),
        "rows": len(expected),
        "columns": 23,
        "moisture_max_abs_difference": diff,
        "terminal_time_difference_s": float(terminal_time_difference),
        "workbook_hash": file_record(workbook),
        "headers": workbook_headers(record),
        "verification_hash": file_record(directory / "verification.json"),
        "delivery_sources": delivery_sources(),
        "table6_hash": file_record(directory / "table6.csv"),
        "radius_input_manifest_hash": file_record(radius_manifest_path),
        "radius_observations_hash": file_record(directory / "radius_observations.csv"),
        "writer": "openpyxl normal mode",
    }
    write_json(directory / "export_verification.json", result)
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--directory", default="results/q4")
    parser.add_argument("--workbook", default="results/result4.xlsx")
    args = parser.parse_args()
    print(check(args.directory, args.workbook))

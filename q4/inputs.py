"""Problem 4 input binding: environment, radius history, and result template."""
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator

from q2.inputs import LongEnvironment, read_config, read_inputs as read_q2_inputs, sha256


@dataclass(frozen=True)
class RadiusHistory:
    observations: np.ndarray
    fixed: bool = False
    interpolation: str = "linear"
    horizon_s: float | None = None

    def __post_init__(self):
        values = np.asarray(self.observations, dtype=float)
        if values.ndim != 2 or values.shape[1] != 2 or not np.all(np.isfinite(values)):
            raise ValueError("Expected finite two-column radius observations")
        if values[0, 0] != 0 or np.any(np.diff(values[:, 0]) <= 0):
            raise ValueError("Radius time axis must start at 0 and be strictly increasing")
        if np.any(values[:, 1] <= 0) or np.any(np.diff(values[:, 1]) > 1e-12):
            raise ValueError("Radius observations must stay positive and nonincreasing")
        if self.interpolation not in {"linear", "pchip"}:
            raise ValueError("Radius interpolation must be linear or pchip")
        horizon = values[-1, 0] if self.horizon_s is None else float(self.horizon_s)
        if not np.isfinite(horizon) or horizon <= 0:
            raise ValueError("Radius horizon must be positive and finite")
        if not self.fixed and horizon > values[-1, 0]:
            raise ValueError("A shrinking radius cannot be extrapolated past observations")
        object.__setattr__(self, "observations", values.copy())
        object.__setattr__(self, "horizon_s", horizon)
        interpolator = None
        if not self.fixed and self.interpolation == "pchip":
            interpolator = PchipInterpolator(values[:, 0], values[:, 1], extrapolate=False)
            probe = np.linspace(0.0, horizon, max(1001, 8 * len(values)))
            radii = np.asarray(interpolator(probe), dtype=float)
            if np.any(~np.isfinite(radii)) or np.any(radii <= 0) or np.any(np.diff(radii) > 1e-12):
                raise ValueError("PCHIP radius must remain finite, positive, and nonincreasing")
        object.__setattr__(self, "_interpolator", interpolator)

    def __call__(self, t):
        query = np.asarray(t, dtype=float)
        if np.any(~np.isfinite(query)) or np.any(query < 0) or np.any(query > self.horizon_s):
            raise ValueError("Radius query outside configured history")
        if self.fixed:
            return np.full_like(query, self.observations[0, 1], dtype=float)
        if self.interpolation == "pchip":
            return np.asarray(self._interpolator(query), dtype=float)
        return np.interp(query, self.observations[:, 0], self.observations[:, 1])

    @property
    def breaks(self):
        values = self.observations[self.observations[:, 0] <= self.horizon_s, 0]
        return values if values[-1] == self.horizon_s else np.r_[values, self.horizon_s]


def read_radius(path, expected_observed_end_s=None):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        rows = list(workbook.active.values)
    finally:
        workbook.close()
    numeric = np.asarray([row[:2] for row in rows[1:] if row[0] is not None and row[1] is not None], dtype=float)
    if numeric.shape != (145, 2):
        raise ValueError("Expected 145 radius rows after the header")
    numeric[:, 1] *= 0.01
    if expected_observed_end_s is not None and numeric[-1, 0] != float(expected_observed_end_s):
        raise ValueError("Unexpected radius observation endpoint")
    return RadiusHistory(numeric)


def read_inputs(
    data_root,
    config,
    mode=None,
    fixed_radius=False,
    interpolation=None,
    radius_offset_cm=0.0,
    horizon_s=None,
):
    q2_config = read_config(config["q2_config"])
    base_environment, metadata = read_q2_inputs(data_root, q2_config, mode=mode)
    selected_horizon = float(horizon_s or q2_config["horizon_s"])
    selected_mode = mode or q2_config["environment"]["default_extension"]
    environment = LongEnvironment(
        base_environment.observations,
        selected_horizon,
        selected_mode,
        q2_config["environment"]["mean_window_start_s"],
    )
    root = Path(data_root)
    paths = {
        "radius": root / "附件" / "附件2.xlsx",
        "template": root / "附件" / "附件3" / "result4.xlsx",
    }
    hashes = {key: sha256(path) for key, path in paths.items()}
    for key, expected in config["input_sha256"].items():
        if hashes[key] != expected:
            raise ValueError(f"Unreviewed Q4 {key} input hash: {hashes[key]}")
    expected_radius_end_s = float(config["radius"]["observed_end_s"])
    radius = read_radius(paths["radius"], expected_radius_end_s)
    offset_m = float(radius_offset_cm) * 0.01
    if offset_m:
        observations = radius.observations.copy()
        observations[1:, 1] += offset_m
        radius = RadiusHistory(observations)
    selected_interpolation = interpolation or config["radius"]["interpolation"]
    if fixed_radius:
        radius = RadiusHistory(
            radius.observations,
            fixed=True,
            interpolation=selected_interpolation,
            horizon_s=selected_horizon,
        )
    else:
        radius = RadiusHistory(
            radius.observations,
            interpolation=selected_interpolation,
            horizon_s=selected_horizon,
        )
    template = load_workbook(paths["template"], data_only=True, read_only=True)
    try:
        if template.sheetnames != ["Sheet1"]:
            raise ValueError("Unexpected result4 template sheets")
        template_a1 = template.active["A1"].value
        template_surface = template.active["F1"].value
    finally:
        template.close()
    metadata["q4_radius"] = {
        "sha256": hashes["radius"],
        "source": str(paths["radius"].relative_to(root)),
        "rows": int(len(radius.observations)),
        "initial_radius_m": float(radius.observations[0, 1]),
        "final_radius_m": float(radius.observations[-1, 1]),
        "fixed": fixed_radius,
        "interpolation": selected_interpolation,
        "record_offset_cm": float(radius_offset_cm),
        "observed_end_s": float(radius.observations[-1, 0]),
        "horizon_s": selected_horizon,
    }
    metadata["environment_extension"] = {
        **metadata["environment_extension"],
        "mode": selected_mode,
        "temperature_C": float(environment.extension[0]),
        "moisture": float(environment.extension[1]),
        "horizon_s": selected_horizon,
    }
    metadata["q4_template"] = {
        "sha256": hashes["template"],
        "source": str(paths["template"].relative_to(root)),
        "A1": template_a1,
        "surface_header": template_surface,
    }
    return environment, radius, metadata


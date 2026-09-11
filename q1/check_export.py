"""Independent openpyxl readback of every delivered value and format."""
import argparse
import csv
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from .archive import write_json
from .export import verified_source,rounded_array,TABLE_TIMES
from .inputs import sha256


def check(directory="results/q1",workbook="results/result1.xlsx"):
    directory = Path(directory)
    data,g,m,v = verified_source(directory)
    wb = load_workbook(workbook,read_only=False,data_only=True)
    if wb.sheetnames != ["温度","水分浓度"]:
        raise ValueError("Workbook sheets differ from contract")
    count = 0
    for sheet,key,table_key in zip(wb.worksheets,["temperature_C","moisture"],["temperature","moisture"]):
        if sheet.max_row != 1801 or sheet.max_column != 22 or sheet["A1"].value != m["inputs"]["template_A1"]:
            raise ValueError("Workbook dimensions/header differ from contract")
        if sheet.freeze_panes != "B2":
            raise ValueError(f"Header/time freeze panes missing: {sheet.freeze_panes}")
        if [c.value for c in sheet[1][1:]] != [j/10 for j in range(21)]:
            raise ValueError("Radius coordinates mismatch")
        expected = rounded_array(data[key][1:,g["output_indices"]])
        actual = []
        for i,row in enumerate(sheet.iter_rows(min_row=2),start=1):
            if row[0].value != i:
                raise ValueError("Time column mismatch")
            for cell in row[1:]:
                if cell.data_type != "n" or cell.number_format != "0.0000":
                    raise ValueError("Result cell is not a numeric four-decimal value")
            actual.append([c.value for c in row[1:]])
        if not np.array_equal(actual,expected):
            raise ValueError("Full workbook value readback mismatch")
        with (directory/f"table_{table_key}.csv").open(encoding="utf-8",newline="") as stream:
            table = np.asarray(list(csv.reader(stream))[1:],dtype=float)
        if not np.array_equal(table[:,0],TABLE_TIMES) or not np.array_equal(table[:,1:],expected[TABLE_TIMES-1][:,[0,5,10,15,20]]):
            raise ValueError("Paper table/Excel inconsistency")
        count += expected.size
    wb.close()
    # Plots read these same CSVs. Check every archived plotted value independently.
    figure_dir = directory/"figure_data"
    response = np.loadtxt(figure_dir/"response.csv",delimiter=",",skiprows=1)
    expected_response = np.c_[data["time_s"],data["temperature_C"][:,0],data["temperature_C"][:,-1],data["temperature_C"]@g["volume_m3"]/g["volume_m3"].sum(),data["moisture"][:,0],data["moisture"][:,-1],data["moisture"]@g["volume_m3"]/g["volume_m3"].sum()]
    if not np.array_equal(response,expected_response):
        raise ValueError("Response figure data mismatch")
    profiles = np.loadtxt(figure_dir/"profiles.csv",delimiter=",",skiprows=1)
    length = len(g["radius_m"])
    if not np.array_equal(profiles[:,0],np.repeat(TABLE_TIMES,length)) or not np.array_equal(profiles[:,1],np.tile(g["radius_m"]*100,len(TABLE_TIMES))):
        raise ValueError("Profile figure coordinates mismatch")
    for col,key in [(2,"temperature_C"),(3,"moisture")]:
        if not np.array_equal(profiles[:,col],data[key][TABLE_TIMES].ravel()):
            raise ValueError("Profile figure data mismatch")
    result = {"passed":True,"numeric_result_cells_checked":count,"workbook_sha256":sha256(workbook),"table_values_match":True,"figure_source_values_match":True,"sheets":["温度","水分浓度"],"range_per_sheet":"A1:V1801","rounding":"Decimal(str(float)); ROUND_HALF_UP; 4 decimal places","original_template_sha256":m["inputs"]["sha256"]["template"]}
    write_json(directory/"export_verification.json",result)
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--directory",default="results/q1")
    parser.add_argument("--workbook",default="results/result1.xlsx")
    args = parser.parse_args()
    print(json.dumps(check(args.directory,args.workbook),ensure_ascii=False))

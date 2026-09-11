"""Read-only source inputs; no extrapolation, guessed unit conversions or edits."""
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def read_config(path="configs/q1.json"):
    return json.loads(Path(path).read_text(encoding="utf-8"))


class Environment:
    def __init__(self, observations):
        a = np.asarray(observations, dtype=float)
        if a.shape != (31, 3) or not np.all(np.isfinite(a)) or not np.array_equal(a[:,0], np.arange(0,1801,60)):
            raise ValueError("Expected 31 finite records at 0:60:1800 seconds")
        if np.any(a[:,2] <= 0):
            raise ValueError("Equivalent moisture drive must be positive")
        self.observations = a.copy()
        self._history_min = np.minimum.accumulate(a[:,1:], axis=0)
        self._history_max = np.maximum.accumulate(a[:,1:], axis=0)

    def __call__(self, t):
        t = np.asarray(t)
        if np.any(~np.isfinite(t)) or np.any(t < 0) or np.any(t > 1800):
            raise ValueError("Q1 environment only defined on [0, 1800] s")
        a = self.observations
        return np.interp(t, a[:,0], a[:,1]), np.interp(t, a[:,0], a[:,2])

    def history_extrema(self, t):
        """Exact extrema of the piecewise linear forcing on [0,t].

        An interior query adds its interpolated endpoint, not the next knot.
        Return Tmin,Tmax,Cmin,Cmax; initial material bounds are added by solver.
        """
        T, C = self(t)  # Also rejects nonfinite/out-of-range queries.
        j = np.searchsorted(self.observations[:,0], t, side="right") - 1
        current = np.stack((T,C), axis=-1)
        low = np.minimum(self._history_min[j], current)
        high = np.maximum(self._history_max[j], current)
        return low[...,0], high[...,0], low[...,1], high[...,1]


def read_inputs(data_root, config):
    root = Path(data_root)
    paths = {"environment": root/"附件"/"附件1.xlsx", "template": root/"附件"/"附件3"/"result1.xlsx", "problem": root/"A题.pdf"}
    hashes = {key: sha256(path) for key,path in paths.items()}
    for key, expected in config["input_sha256"].items():
        if hashes[key] != expected:
            raise ValueError(f"Unreviewed {key} input hash: {hashes[key]}")
    wb = load_workbook(paths["environment"], data_only=True, read_only=True)
    try:
        rows = list(wb["Sheet1"].values)
    finally:
        wb.close()
    if rows[0][:3] != ("时间", "温度", "水分浓度"):
        raise ValueError("Unexpected environment headers")
    all_data = np.asarray([row[:3] for row in rows[1:]], dtype=float)
    if all_data.shape != (241,3) or not np.all(np.isfinite(all_data)) or not np.array_equal(all_data[:,0], np.arange(0,14401,60)):
        raise ValueError("Unexpected source time axis or missing data")
    template = load_workbook(paths["template"], read_only=True)
    try:
        if template.sheetnames != ["温度", "水分浓度"]:
            raise ValueError("Unexpected template sheets")
        header = template.worksheets[0]["A1"].value
    finally:
        template.close()
    return Environment(all_data[:31]), {"sha256": hashes, "source_names": {k:str(v.relative_to(root)) for k,v in paths.items()}, "template_A1": header}

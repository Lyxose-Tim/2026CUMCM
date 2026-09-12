import pytest
from openpyxl import Workbook

from common.portable_audit import audit
from q2.check_export import frozen_sheet_names


def test_portable_audit_orchestrator_does_not_require_git(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    result = audit({"sentinel": lambda: {"checked": True}})
    assert result["passed"] is True
    assert result["git_required"] is False
    assert result["sections"] == {"sentinel": {"checked": True}}


def test_portable_audit_propagates_failed_section():
    def fail():
        raise ValueError("changed artifact")

    with pytest.raises(ValueError, match="changed artifact"):
        audit({"sentinel": fail})


def test_freeze_panes_are_read_from_streaming_xlsx_metadata(tmp_path):
    path = tmp_path / "stream.xlsx"
    workbook = Workbook(write_only=True)
    first = workbook.create_sheet("温度")
    first.freeze_panes = "B2"
    first.append(["time", 0.0])
    second = workbook.create_sheet("水分浓度")
    second.freeze_panes = "B2"
    second.append(["time", 0.0])
    workbook.save(path)
    assert frozen_sheet_names(path) == ["温度", "水分浓度"]

import numpy as np
import pytest

from q3.check_export import compare_numeric
from q3.export import rounded


@pytest.mark.parametrize("bad", [None,np.nan,np.inf,-np.inf,"0.15",True])
def test_missing_or_invalid_result_cannot_mask_second_error(bad):
    with pytest.raises(ValueError):
        compare_numeric([[bad,999.0]],np.array([[0.1,0.2]]))


def test_incorrect_finite_result_is_rejected():
    with pytest.raises(ValueError):
        compare_numeric([[0.1,0.25]],np.array([[0.1,0.2]]))


def test_half_up_rounding_does_not_determine_dryness():
    values=np.array([0.1499999,0.1500001,0.12345])
    np.testing.assert_array_equal(rounded(values),[0.15,0.15,0.1235])
    assert (values[0]<0.15) and not (values[1]<0.15)


def test_failed_export_blocks_reports(tmp_path):
    import json
    from q3.reports import require_export
    (tmp_path/"export_verification.json").write_text(json.dumps({"passed":False}),encoding="utf-8")
    with pytest.raises(ValueError,match="failed"):
        require_export(tmp_path,tmp_path/"nonexistent.xlsx")


def test_changed_workbook_blocks_reports(tmp_path):
    import json
    from q3.reports import require_export
    (tmp_path/"export_verification.json").write_text(json.dumps({"passed":True,"workbook_sha256":"wrong"}),encoding="utf-8")
    (tmp_path/"result.xlsx").write_bytes(b"a changed workbook")
    with pytest.raises(ValueError,match="different workbook"):
        require_export(tmp_path,tmp_path/"result.xlsx")


def test_dimensionless_artifact_tool_sheet_requires_normal_load(tmp_path):
    # A minimal XLSX written through Artifact Tool may omit worksheet
    # dimension metadata: read-only max_row is None, while normal load
    # correctly discovers the populated cells used by the verifier.
    from openpyxl import load_workbook
    path="results/result3.xlsx"
    ro=load_workbook(path,read_only=True,data_only=True)
    try:
        assert ro.active.max_row is None
    finally:
        ro.close()
    normal=load_workbook(path,read_only=False,data_only=True)
    try:
        assert normal.active.max_row > 3000 and normal.active.max_column == 22
    finally:
        normal.close()

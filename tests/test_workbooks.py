import numpy as np
from openpyxl import load_workbook

from common.workbooks import write_dense_workbook, write_stream_workbook


def sheet_spec():
    return {
        "name": "Sheet1",
        "header": ["time", 0.0, "surface"],
        "times": np.array([1.0, 2.0]),
        "values": np.array([[1.23445, np.nan], [2.0, 3.0]]),
    }


def assert_plain_python_workbook(path):
    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        sheet = workbook["Sheet1"]
        assert sheet.freeze_panes == "B2"
        assert sheet.sheet_view.showGridLines is False
        assert sheet.max_row == 3 and sheet.max_column == 3
        assert sheet["B2"].value == 1.2345
        assert sheet["B2"].data_type == "n"
        assert sheet["B2"].number_format == "0.0000"
        assert sheet["C2"].value is None
    finally:
        workbook.close()


def test_dense_openpyxl_writer_has_numeric_cells_and_domain_blanks(tmp_path):
    path = tmp_path / "dense.xlsx"
    write_dense_workbook(path, [sheet_spec()])
    assert_plain_python_workbook(path)


def test_streaming_openpyxl_writer_has_same_contract(tmp_path):
    path = tmp_path / "stream.xlsx"
    write_stream_workbook(path, [sheet_spec()])
    assert_plain_python_workbook(path)

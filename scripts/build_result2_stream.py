"""Stream the verified Q2 payload into the final large XLSX workbook."""
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill


def sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def cell(sheet, value, number_format, font=None, fill=None):
    result = WriteOnlyCell(sheet, value=value)
    result.number_format = number_format
    if font is not None:
        result.font = font
    if fill is not None:
        result.fill = fill
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--payload", default=".scratch/q2_workbook/payload.json")
    parser.add_argument("--output", default="results/result2.xlsx")
    args = parser.parse_args()
    payload = json.loads(Path(args.payload).read_text(encoding="utf-8"))
    for key in ("verification", "archive_manifest"):
        if sha256(payload[f"{key}_file"]) != payload[f"{key}_sha256"]:
            raise ValueError(f"{key} hash mismatch")

    workbook = Workbook(write_only=True)
    normal = Font(name="Arial", size=10, color="1F2937")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    names = ["\u6e29\u5ea6", "\u6c34\u5206\u6d53\u5ea6"]
    sheets = {}
    header = [payload["template_A1"], *[index / 10 for index in range(21)]]
    for name in names:
        sheet = workbook.create_sheet(name)
        sheet.freeze_panes = "B2"
        sheet.sheet_view.showGridLines = False
        sheet.column_dimensions["A"].width = 12
        for letter in "BCDEFGHIJKLMNOPQRSTUV":
            sheet.column_dimensions[letter].width = 11
        sheet.append([cell(sheet, value, "General", header_font, header_fill) for value in header])
        sheets[name] = sheet

    counts = {name: 0 for name in names}
    for chunk in payload["chunks"]:
        path = Path(chunk["file"])
        if sha256(path) != chunk["sha256"]:
            raise ValueError(f"Chunk hash mismatch: {path}")
        matrix = json.loads(path.read_text(encoding="utf-8"))
        if len(matrix) != chunk["rows"]:
            raise ValueError(f"Chunk row mismatch: {path}")
        sheet = sheets[chunk["sheet"]]
        for row in matrix:
            if len(row) != 22:
                raise ValueError(f"Chunk column mismatch: {path}")
            sheet.append([
                cell(sheet, int(row[0]), "0", normal),
                *[cell(sheet, float(value), "0.0000", normal) for value in row[1:]],
            ])
        counts[chunk["sheet"]] += len(matrix)
    if any(count != 259200 for count in counts.values()):
        raise ValueError(f"Incomplete formal rows: {counts}")
    destination = Path(args.output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    print(json.dumps({"output": str(destination), "bytes": destination.stat().st_size, "rows": counts}, ensure_ascii=False))


if __name__ == "__main__":
    main()
